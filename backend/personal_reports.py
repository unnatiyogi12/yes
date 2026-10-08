"""AURA personal progress reports: GitHub, LeetCode, Gmail and Calendar.

Google services require a one-time OAuth consent using credentials.json.
LeetCode uses its public profile GraphQL endpoint; no API key is required.
"""
from __future__ import annotations

import json
import re
import subprocess
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import requests

BASE_DIR = Path(__file__).resolve().parent
BACKEND_DIR = BASE_DIR
REPORT_DIR = BACKEND_DIR / "reports"
CONFIG_FILE = BACKEND_DIR / "aura_personal_config.json"
GOOGLE_CREDENTIALS = BACKEND_DIR / "credentials.json"
GOOGLE_TOKEN = BACKEND_DIR / ".aura_google_token.json"

LEETCODE_URL = "https://leetcode.com/graphql"


def _run(cmd: list[str], cwd: str | None = None, timeout: int = 60) -> tuple[int, str, str]:
    p = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, shell=False)
    return p.returncode, p.stdout.strip(), p.stderr.strip()


def load_config() -> dict[str, Any]:
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def save_config(data: dict[str, Any]) -> None:
    CONFIG_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")


def set_leetcode_username(username: str) -> str:
    username = username.strip().lstrip("@").rstrip("/").split("/")[-1]
    if not re.fullmatch(r"[A-Za-z0-9_-]+", username):
        raise ValueError("Invalid LeetCode username.")
    cfg = load_config()
    cfg["leetcode_username"] = username
    save_config(cfg)
    return username


def get_leetcode_username() -> str | None:
    return load_config().get("leetcode_username")


def fetch_leetcode(username: str | None = None) -> dict[str, Any]:
    username = username or get_leetcode_username()
    if not username:
        raise RuntimeError("LeetCode username is not configured. Use: set leetcode username YOUR_USERNAME")

    query = """
    query userData($username: String!) {
      matchedUser(username: $username) {
        username
        profile { realName userAvatar ranking reputation }
        submitStatsGlobal { acSubmissionNum { difficulty count submissions } totalSubmissionNum { difficulty count submissions } }
        problemsSolvedBeatsStats { difficulty percentage }
      }
      recentAcSubmissionList(username: $username, limit: 20) {
        title titleSlug timestamp lang
      }
    }
    """
    r = requests.post(LEETCODE_URL, json={"query": query, "variables": {"username": username}}, timeout=30)
    r.raise_for_status()
    data = r.json()
    if data.get("errors"):
        raise RuntimeError(data["errors"][0].get("message", "LeetCode request failed."))
    user = data.get("data", {}).get("matchedUser")
    if not user:
        raise RuntimeError(f"LeetCode profile not found for '{username}'.")
    stats = user.get("submitStatsGlobal", {}) or {}
    solved = {x["difficulty"]: x for x in stats.get("acSubmissionNum", [])}
    total = {x["difficulty"]: x for x in stats.get("totalSubmissionNum", [])}
    recent = data.get("data", {}).get("recentAcSubmissionList") or []
    return {
        "username": username,
        "profile": user.get("profile") or {},
        "solved": {k: int(v.get("count", 0)) for k, v in solved.items()},
        "submissions": {k: int(v.get("submissions", 0)) for k, v in total.items()},
        "recent": recent,
    }


def github_report(days: int = 30) -> dict[str, Any]:
    code, out, err = _run(["gh", "auth", "status"], timeout=20)
    if code != 0:
        raise RuntimeError("GitHub CLI is not authenticated. Run `gh auth login`.")
    code, out, err = _run(["gh", "api", "user"], timeout=30)
    if code != 0:
        raise RuntimeError(err or "Could not read GitHub profile.")
    user = json.loads(out)
    login = user.get("login")

    code, out, err = _run(["gh", "repo", "list", login, "--limit", "100", "--json", "name,isPrivate,updatedAt,pushedAt,url,description,primaryLanguage"], timeout=60)
    if code != 0:
        raise RuntimeError(err or "Could not list GitHub repositories.")
    repos = json.loads(out or "[]")
    since = datetime.now(timezone.utc) - timedelta(days=days)
    commits = []
    for repo in repos:
        name = repo.get("name")
        code, out, err = _run(["gh", "api", f"repos/{login}/{name}/commits", "--method", "GET", "-f", f"since={since.isoformat().replace('+00:00','Z')}", "-f", "per_page=100"], timeout=30)
        if code != 0:
            continue
        try:
            rows = json.loads(out or "[]")
        except Exception:
            rows = []
        for row in rows:
            commit = row.get("commit", {})
            commits.append({"repo": name, "sha": row.get("sha", "")[:7], "message": (commit.get("message") or "").splitlines()[0], "date": commit.get("author", {}).get("date", "")})

    return {"profile": user, "repos": repos, "commits": commits, "period_days": days}


def _google_credentials():
    try:
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from google.auth.transport.requests import Request
    except ImportError as exc:
        raise RuntimeError("Google integration libraries are missing. Install: google-api-python-client google-auth-httplib2 google-auth-oauthlib") from exc

    scopes = [
        "https://www.googleapis.com/auth/gmail.readonly",
        "https://www.googleapis.com/auth/calendar.readonly",
        "https://www.googleapis.com/auth/drive.readonly",
    ]
    creds = None
    if GOOGLE_TOKEN.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(GOOGLE_TOKEN), scopes)
        except Exception:
            creds = None
    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
    if not creds or not creds.valid:
        if not GOOGLE_CREDENTIALS.exists():
            raise RuntimeError(f"Google OAuth setup needed: put your downloaded credentials.json in {GOOGLE_CREDENTIALS}")
        flow = InstalledAppFlow.from_client_secrets_file(str(GOOGLE_CREDENTIALS), scopes)
        creds = flow.run_local_server(port=0)
        GOOGLE_TOKEN.write_text(creds.to_json(), encoding="utf-8")
    return creds


def google_report(days: int = 30) -> dict[str, Any]:
    from googleapiclient.discovery import build
    creds = _google_credentials()
    gmail = build("gmail", "v1", credentials=creds, cache_discovery=False)
    calendar = build("calendar", "v3", credentials=creds, cache_discovery=False)

    after = int((datetime.now(timezone.utc) - timedelta(days=days)).timestamp())
    g = gmail.users().messages().list(userId="me", q=f"after:{after}", maxResults=50).execute()
    messages = []
    for item in g.get("messages", [])[:50]:
        msg = gmail.users().messages().get(userId="me", id=item["id"], format="metadata", metadataHeaders=["Subject", "From", "Date"]).execute()
        headers = {h["name"]: h["value"] for h in msg.get("payload", {}).get("headers", [])}
        messages.append({"id": item["id"], "subject": headers.get("Subject", "(no subject)"), "from": headers.get("From", ""), "date": headers.get("Date", ""), "snippet": msg.get("snippet", "")})

    now = datetime.now(timezone.utc)
    events = calendar.events().list(calendarId="primary", timeMin=(now - timedelta(days=days)).isoformat(), timeMax=(now + timedelta(days=days)).isoformat(), singleEvents=True, orderBy="startTime", maxResults=100).execute()
    event_rows = []
    for e in events.get("items", []):
        start = e.get("start", {}).get("dateTime") or e.get("start", {}).get("date", "")
        event_rows.append({"summary": e.get("summary", "(untitled)"), "start": start, "status": e.get("status", "")})
    return {"emails": messages, "events": event_rows, "period_days": days}


def collect_progress(days: int = 30) -> dict[str, Any]:
    result: dict[str, Any] = {"generated_at": datetime.now().isoformat(timespec="seconds"), "period_days": days}
    for key, fn in (("github", lambda: github_report(days)), ("leetcode", lambda: fetch_leetcode()), ("google", lambda: google_report(days))):
        try:
            result[key] = {"ok": True, "data": fn()}
        except Exception as exc:
            result[key] = {"ok": False, "error": str(exc)}
    return result


def _add_bullets(doc, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def generate_docx_report(data: dict[str, Any], days: int = 30) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is missing. Install it with: pip install python-docx") from exc

    REPORT_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = REPORT_DIR / f"AURA_Recent_Progress_Report_{stamp}.docx"
    doc = Document()
    title = doc.add_heading("AURA AI — Recent Progress Report", 0)
    doc.add_paragraph(f"Reporting period: last {days} days")
    doc.add_paragraph(f"Generated: {data['generated_at']}")

    gh = data.get("github", {})
    doc.add_heading("GitHub", 1)
    if gh.get("ok"):
        d = gh["data"]
        _add_bullets(doc, [f"Account: {d['profile'].get('login')}", f"Repositories: {len(d['repos'])}", f"Commits in period: {len(d['commits'])}"])
        if d["commits"]:
            doc.add_paragraph("Recent commits:")
            _add_bullets(doc, [f"{c['repo']}: {c['message']} ({c['date'][:10]})" for c in d["commits"][:15]])
    else:
        doc.add_paragraph(f"Unavailable: {gh.get('error')}")

    lc = data.get("leetcode", {})
    doc.add_heading("LeetCode", 1)
    if lc.get("ok"):
        d = lc["data"]
        solved = d.get("solved", {})
        _add_bullets(doc, [f"Username: {d['username']}", f"Total solved: {sum(solved.values())}", f"Easy: {solved.get('Easy', 0)} | Medium: {solved.get('Medium', 0)} | Hard: {solved.get('Hard', 0)}", f"Ranking: {d['profile'].get('ranking', 'N/A')}"])
        if d.get("recent"):
            doc.add_paragraph("Recent accepted problems:")
            _add_bullets(doc, [f"{x.get('title')} — {x.get('lang')}" for x in d["recent"][:10]])
    else:
        doc.add_paragraph(f"Unavailable: {lc.get('error')}")

    google = data.get("google", {})
    doc.add_heading("Gmail & Calendar", 1)
    if google.get("ok"):
        d = google["data"]
        _add_bullets(doc, [f"Emails found in period: {len(d['emails'])}", f"Calendar events in window: {len(d['events'])}"])
        if d["emails"]:
            doc.add_paragraph("Recent emails:")
            _add_bullets(doc, [f"{e['subject']} — {e['from']}" for e in d["emails"][:10]])
        if d["events"]:
            doc.add_paragraph("Calendar activity:")
            _add_bullets(doc, [f"{e['summary']} — {e['start']}" for e in d["events"][:10]])
    else:
        doc.add_paragraph(f"Unavailable: {google.get('error')}")

    doc.add_heading("AURA Summary", 1)
    summaries = []
    if gh.get("ok"): summaries.append(f"GitHub recorded {len(gh['data']['commits'])} commits across {len(gh['data']['repos'])} repositories.")
    if lc.get("ok"): summaries.append(f"LeetCode shows {sum(lc['data'].get('solved', {}).values())} total solved problems for {lc['data']['username']}.")
    if google.get("ok"): summaries.append(f"Google services returned {len(google['data']['emails'])} emails and {len(google['data']['events'])} calendar events for the reporting window.")
    _add_bullets(doc, summaries or ["No connected data sources were available yet."])

    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Aptos"
            style.font.size = Pt(10)
    doc.save(path)
    return path
